import os
os.environ["TESSDATA_PREFIX"] = "/usr/share/tesseract-ocr/4.00/tessdata"
import PyPDF2
import io
import frappe
import pytesseract
import numpy as np
from PIL import Image
from gdoc.gdoc.preprocess_image import extract_scanned_pdf,extract_image,process_pil_image
from gdoc.gdoc.digital_pdf_extract import extract_digital_pdf_elements 


def _is_useful_image(pil_image: Image.Image) -> bool:
    """Skip tiny images like icons, logos."""
    width, height = pil_image.size
    return width >= 100 and height >= 100


def _extract_from_image_bytes(image_bytes: bytes, lang: str = "eng+ara") -> str:
    """
    Full extraction from raw image bytes:
    preprocess → table extraction → full text extraction
    """
    try:
        pil_image = Image.open(io.BytesIO(image_bytes))

        if not _is_useful_image(pil_image):
            return ""

        res = process_pil_image(pil_image,lang)
        return res

    except Exception as e:
        frappe.log_error(str(e)[:140], "_extract_from_image_bytes failed")
        return ""


def is_scanned_pdf(pdf_path):
    with open(pdf_path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        chars = sum(len((p.extract_text() or "").strip()) for p in reader.pages)
    return chars < 50 * len(reader.pages)   # avg <50 chars/page => scanned


def _extract_from_pdf(full_file_path: str, lang: str = "eng+ara") -> str:
    try:
        if is_scanned_pdf(full_file_path):
            return extract_scanned_pdf(full_file_path,lang)
        else:
            return extract_digital_pdf_elements(full_file_path) 
    except Exception as e:
        frappe.log_error(str(e)[:140], "_extract_from_pdf failed")
        return ""


def _extract_from_docx(full_file_path: str, lang: str = "eng+ara") -> str:
    try:
        import docx
        import pandas as pd

        doc   = docx.Document(full_file_path)
        parts = []

        for element in doc.element.body:
            tag = element.tag.split("}")[-1]

            if tag == "p":
                para = docx.text.paragraph.Paragraph(element, doc)
                text = para.text.strip()
                if text:
                    parts.append(text)

            elif tag == "tbl":
                table = docx.table.Table(element, doc)
                rows  = [
                    [cell.text.strip() for cell in row.cells]
                    for row in table.rows
                ]
                if rows:
                    df = pd.DataFrame(rows[1:], columns=rows[0])
                    parts.append(df.to_markdown(index=False))

        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                try:
                    content = _extract_from_image_bytes(
                        rel.target_part.blob, lang
                    )
                    if content:
                        parts.append(content)
                except Exception:
                    pass

        return "\n\n".join(parts)

    except Exception as e:
        frappe.log_error(str(e)[:140], "_extract_from_docx failed")
        return ""


def _extract_from_xlsx(full_file_path: str, lang: str = "eng+ara") -> str:
    try:
        import openpyxl
        import pandas as pd

        wb    = openpyxl.load_workbook(full_file_path, data_only=True)
        parts = []

        for sheet_name in wb.sheetnames:
            ws          = wb[sheet_name]
            sheet_parts = [f"## Sheet: {sheet_name}"]

            data = [
                [str(cell) if cell is not None else "" for cell in row]
                for row in ws.iter_rows(values_only=True)
                if any(cell is not None for cell in row)
            ]
            if data:
                df = pd.DataFrame(data[1:], columns=data[0]) \
                    if len(data) > 1 else pd.DataFrame(data)
                sheet_parts.append(df.to_markdown(index=False))

            for image in ws._images:
                try:
                    content = _extract_from_image_bytes(
                        image._data(), lang
                    )
                    if content:
                        sheet_parts.append(content)
                except Exception:
                    pass

            parts.append("\n\n".join(sheet_parts))

        return "\n\n".join(parts)

    except Exception as e:
        frappe.log_error(str(e)[:140], "_extract_from_xlsx failed")
        return ""


def _extract_from_pptx(full_file_path: str, lang: str = "eng+ara") -> str:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        import pandas as pd

        prs   = Presentation(full_file_path)
        parts = []

        for slide_no, slide in enumerate(prs.slides, 1):
            slide_parts = [f"## Slide {slide_no}"]

            for shape in slide.shapes:

                if shape.has_text_frame:
                    text = "\n".join(
                        para.text.strip()
                        for para in shape.text_frame.paragraphs
                        if para.text.strip()
                    )
                    if text:
                        slide_parts.append(text)

                if shape.has_table:
                    rows = [
                        [cell.text.strip() for cell in row.cells]
                        for row in shape.table.rows
                    ]
                    if rows:
                        df = pd.DataFrame(rows[1:], columns=rows[0])
                        slide_parts.append(df.to_markdown(index=False))

                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        content = _extract_from_image_bytes(
                            shape.image.blob, lang
                        )
                        if content:
                            slide_parts.append(content)
                    except Exception:
                        pass

            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_parts.append(f"Notes: {notes}")

            parts.append("\n\n".join(slide_parts))

        return "\n\n".join(parts)

    except Exception as e:
        frappe.log_error(str(e)[:140], "_extract_from_pptx failed")
        return ""


def _extract_from_office(full_file_path: str, lang: str = "eng+ara") -> str:
    file_ext = os.path.splitext(full_file_path)[1]
    ext = file_ext.lower()
    if ext in [".docx", ".doc"]:
        return _extract_from_docx(full_file_path, lang)
    elif ext in [".xlsx", ".xls"]:
        return _extract_from_xlsx(full_file_path, lang)
    elif ext in [".pptx", ".ppt"]:
        return _extract_from_pptx(full_file_path, lang)
    else:
        frappe.log_error(
            f"Unsupported office format: {ext}",
            "_extract_from_office: unsupported"
        )
        return ""


def _extract_from_text_file(full_file_path: str) -> str:
    try:
        with open(full_file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()
    except Exception as e:
        frappe.log_error(str(e)[:140], "_extract_from_text_file failed")
        return ""


def extract_text_from_file(full_file_path: str, file_ext: str, lang: str = "eng+ara") -> str:
    """
    Production-grade local extraction router.
    All OCR done locally by Tesseract — no cloud, no data leaves server.

    PDF          → fitz (text/tables) + pypdfium2 render + two-pass Tesseract
    Images       → preprocess + two-pass Tesseract + block-level detection
    Word/Excel   → dedicated libraries + Tesseract for embedded images
    Text files   → direct read
    """

    PDF_FORMATS    = [".pdf"]
    IMAGE_FORMATS  = [".jpg", ".jpeg", ".png", ".tiff",
                      ".tif", ".bmp", ".webp", ".gif"]
    OFFICE_FORMATS = [".docx", ".xlsx", ".pptx"]
    TEXT_FORMATS   = [".txt", ".csv", ".json", ".xml",
                      ".html", ".htm", ".md"]

    ext = file_ext.lower()

    if ext in PDF_FORMATS:
        return _extract_from_pdf(full_file_path, lang)

    elif ext in IMAGE_FORMATS:
        return extract_image(full_file_path,lang)

    elif ext in OFFICE_FORMATS:
        return _extract_from_office(full_file_path, lang)

    elif ext in TEXT_FORMATS:
        return _extract_from_text_file(full_file_path)

    else:
        frappe.log_error(
            f"Unsupported file type: {ext}",
            "extract_text_from_file: unsupported format"
        )
        return ""


if _name__ == "__main__" :
    extract_text_from_file()
